import asyncio
import base64
import hashlib
import io
import json
import os
import re
import urllib.parse
import urllib.request
from pathlib import Path

TOKENBOM_API_URL = "https://tokenbom.com/v1/chat/completions"
TOKENBOM_MODEL = "gpt-5.6-luna"
GITHUB_API = "https://api.github.com"
GITHUB_REPO = os.getenv("EFIN_KB_REPO", "Goyardvpn/efin-knowledge-base")
GITHUB_BRANCH = os.getenv("EFIN_KB_BRANCH", "main")
CACHE_DIR = Path("data/github_knowledge")
CACHE_DIR.mkdir(parents=True, exist_ok=True)

SYSTEM_PROMPT = """Ты — Efin AI, внутренний помощник компании EFIN.
Отвечай на русском языке, понятно и по делу.
Главный источник — GitHub-база знаний EFIN.
Используй только материалы выбранного проекта. Не смешивай проекты.
Не выдумывай внутренние правила, тарифы, выплаты и инструкции.
Если пользователь просит фото, изображение, скриншот или пример — используй только реально подходящее изображение из базы.
Если подходящее изображение есть, добавь в конце ответа ровно [SEND_IMAGES].
Если подходящего изображения нет, [SEND_IMAGES] не добавляй.
Если информации нет в базе, прямо скажи об этом.
"""

_PROJECT_CACHE = {}
_FILE_CACHE = {}


def _headers():
    h = {"Accept": "application/vnd.github+json", "User-Agent": "Efin-Money-Bot"}
    token = os.getenv("GITHUB_TOKEN")
    if token:
        h["Authorization"] = f"Bearer {token}"
    return h


def _get_json(url, timeout=45):
    req = urllib.request.Request(url, headers=_headers())
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


def _get_bytes(url, timeout=180):
    req = urllib.request.Request(url, headers=_headers())
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def _decode(data):
    for enc in ("utf-8", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def _scan(path=""):
    encoded = urllib.parse.quote(path, safe="/")
    url = f"{GITHUB_API}/repos/{GITHUB_REPO}/contents/{encoded}?ref={urllib.parse.quote(GITHUB_BRANCH)}"
    return _get_json(url)


def _repo_projects():
    projects = {}
    root_files = []

    def walk(path, out):
        for item in _scan(path):
            if item["type"] == "file":
                out.append(item)
            elif item["type"] == "dir":
                walk(item["path"], out)

    for item in _scan(""):
        if item["type"] == "file":
            root_files.append(item)
        elif item["type"] == "dir":
            files = []
            walk(item["path"], files)
            if files:
                projects[item["name"]] = files

    # Пока МТС лежит в корне репозитория. Когда появятся папки проектов,
    # каждая папка автоматически станет отдельным проектом.
    if root_files:
        projects.setdefault("МБ ( МТС БАНК )", []).extend(root_files)

    return projects


def _projects():
    global _PROJECT_CACHE
    try:
        _PROJECT_CACHE = _repo_projects()
    except Exception as exc:
        if not _PROJECT_CACHE:
            raise
        print(f"[Efin KB] GitHub недоступен, использую предыдущий кэш: {exc}")
    return _PROJECT_CACHE


def _choose_project(projects, query):
    if not projects:
        return None
    q = query.lower()
    aliases = {
        "мб ( мтс банк )": ("мтс", "мтс банк", "мб"),
        "мтс банк": ("мтс", "мтс банк", "мб"),
    }
    ranked = []
    for name in projects:
        score = 0
        for alias in aliases.get(name.lower(), (name.lower(),)):
            if alias in q:
                score += 30 + len(alias)
        score += sum(2 for w in re.findall(r"[а-яa-z0-9ё-]{3,}", name.lower()) if w in q)
        ranked.append((score, name))
    ranked.sort(reverse=True)
    if len(projects) == 1:
        return ranked[0][1]
    return ranked[0][1] if ranked[0][0] else None


def _cached_file(meta):
    path = CACHE_DIR / f"{meta['sha']}_{meta['name']}"
    if path.exists():
        return path.read_bytes()
    data = _get_bytes(meta["download_url"])
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return data


def _image_path(project, meta, index, data, ext):
    directory = CACHE_DIR / "images" / hashlib.sha1(project.encode()).hexdigest()[:12]
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{meta['sha'][:10]}_{index:03d}{ext}"
    if not path.exists():
        path.write_bytes(data)
    return str(path)


def _extract(meta, project):
    key = (project, meta["sha"])
    if key in _FILE_CACHE:
        return _FILE_CACHE[key]
    data = _cached_file(meta)
    name = meta["name"]
    low = name.lower()
    text = ""
    images = []

    if low.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        parts = []
        n = 0
        for page_no, page in enumerate(reader.pages, 1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                parts.append(f"Страница {page_no}:\n{page_text}")
            try:
                for image in page.images:
                    n += 1
                    ext = Path(image.name or ".png").suffix.lower() or ".png"
                    images.append({"path": _image_path(project, meta, n, image.data, ext), "name": f"{name} — страница {page_no}, изображение {n}", "context": page_text, "description": ""})
            except Exception as exc:
                print(f"[Efin KB] PDF image error {name}: {exc}")
        text = "\n\n".join(parts)
    elif low.endswith(".pptx"):
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        n = 0
        for slide_no, slide in enumerate(prs.slides, 1):
            slide_text = []
            for shape in slide.shapes:
                try:
                    if getattr(shape, "has_text_frame", False) and shape.text.strip():
                        slide_text.append(shape.text.strip())
                except Exception:
                    pass
            context = "\n".join(slide_text)
            if context:
                parts.append(f"Слайд {slide_no}:\n{context}")
            for shape in slide.shapes:
                if getattr(shape, "shape_type", None) == 13:
                    try:
                        n += 1
                        ext = "." + (shape.image.ext or "png").lower().lstrip(".")
                        images.append({"path": _image_path(project, meta, n, shape.image.blob, ext), "name": f"{name} — слайд {slide_no}, изображение {n}", "context": context, "description": ""})
                    except Exception as exc:
                        print(f"[Efin KB] PPTX image error {name}: {exc}")
        text = "\n\n".join(parts)
    elif low.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = "\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
    elif low.endswith((".txt", ".md")):
        text = _decode(data)
    elif low.endswith((".jpg", ".jpeg", ".png", ".webp")):
        ext = Path(low).suffix
        text = "Изображение из базы знаний EFIN."
        images = [{"path": _image_path(project, meta, 1, data, ext), "name": name, "context": "", "description": ""}]
    else:
        text = _decode(data)

    result = {"name": name, "text": text.strip(), "images": images}
    _FILE_CACHE[key] = result
    print(f"[Efin KB] Готово: {project} / {name} | текст={len(result['text'])} | фото={len(images)}")
    return result


def _words(text):
    return set(re.findall(r"[а-яa-z0-9ё-]{3,}", (text or "").lower()))


def _relevant(text, query, limit=12000):
    parts = [p.strip() for p in re.split(r"\n\s*\n|\n", text or "") if p.strip()]
    words = _words(query)
    scored = []
    for i, part in enumerate(parts):
        score = sum(1 for w in words if w in part.lower())
        scored.append((score, i, part))
    scored.sort(key=lambda x: (-x[0], x[1]))
    selected = [x for x in scored if x[0] > 0][:40] or scored[:12]
    out, total = [], 0
    for _, i, part in selected:
        block = f"Фрагмент {i + 1}:\n{part}"
        if total + len(block) > limit:
            break
        out.append(block)
        total += len(block)
    return "\n\n".join(out)


def _visual(query):
    q = query.lower()
    return any(x in q for x in ("фото", "фотограф", "сфот", "изображ", "картин", "скрин", "пример", "покажи", "пришли", "как выглядит", "как правильно", "снимок"))


def _image_score(img, query):
    words = _words(query)
    hay = " ".join((img.get("name", ""), img.get("context", ""), img.get("description", ""))).lower()
    score = sum(1 + (2 if len(w) >= 6 else 0) for w in words if w in hay)
    if any(w in words for w in ("клиент", "клиента", "клиентом")) and any(w in hay for w in ("клиент", "фото", "фотограф")):
        score += 6
    return score


def _candidates(images, query, limit=6):
    ranked = sorted(((_image_score(x, query), i, x) for i, x in enumerate(images)), key=lambda x: (-x[0], x[1]))
    return [x for score, _, x in ranked[:limit] if score > 0]


def _data_url(path):
    mime = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}.get(Path(path).suffix.lower(), "image/jpeg")
    return f"data:{mime};base64,{base64.b64encode(Path(path).read_bytes()).decode()}"


def _tokenbom(payload, key, timeout=90):
    req = urllib.request.Request(TOKENBOM_API_URL, data=json.dumps(payload).encode(), headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(req, timeout=timeout) as r:
        data = json.loads(r.read().decode())
    return str(data["choices"][0]["message"]["content"] or "").strip()


async def _load_project(project, metas):
    results = await asyncio.gather(*[asyncio.to_thread(_extract, meta, project) for meta in metas], return_exceptions=True)
    texts, images = [], []
    for meta, result in zip(metas, results):
        if isinstance(result, Exception):
            print(f"[Efin KB] Ошибка {meta['name']}: {result}")
            continue
        texts.append(f"=== ФАЙЛ: {result['name']} ===\n{result['text']}")
        images.extend(result["images"])
    return "\n\n".join(texts), images


async def ask_efin_ai(bot, user_text):
    key = os.getenv("TOKENBOOM_API_KEY")
    if not key:
        raise RuntimeError("Не задана переменная TOKENBOOM_API_KEY")

    projects = _projects()
    project = _choose_project(projects, user_text)
    if not project:
        names = ", ".join(projects) if projects else "нет проектов"
        return {"answer": f"Уточни, по какому проекту вопрос. Доступные проекты: {names}", "images": []}

    text, images = await _load_project(project, projects[project])
    context = _relevant(text, user_text)
    selected = _candidates(images, user_text, 6) if _visual(user_text) else []

    # Если имя файла не помогло, проверяем небольшую выборку vision.
    if _visual(user_text) and not selected and images:
        probe = images[:6]
        try:
            content = [{"type": "text", "text": "Опиши изображения. Отдельно укажи, есть ли человек/клиент и является ли фото примером правильного фото клиента. Формат IMAGE N: описание."}]
            for i, img in enumerate(probe, 1):
                content += [{"type": "text", "text": f"IMAGE {i}: {img['name']}\n{img.get('context','')[:1200]}"}, {"type": "image_url", "image_url": {"url": _data_url(img['path'])}}]
            vision = await asyncio.to_thread(_tokenbom, {"model": TOKENBOM_MODEL, "messages":[{"role":"user","content":content}], "max_tokens":700}, key, 75)
            for line in vision.splitlines():
                m = re.match(r"IMAGE\s+(\d+)\s*:\s*(.*)", line.strip(), re.I)
                if m and 1 <= int(m.group(1)) <= len(probe):
                    probe[int(m.group(1))-1]["description"] = m.group(2)
            selected = _candidates(images, user_text, 6)
        except Exception as exc:
            print(f"[Efin KB] Vision error: {exc}")

    print(f"[Efin AI] GitHub KB: проект={project}; файлов={len(projects[project])}; фото={len(selected)}")
    system = SYSTEM_PROMPT + f"\n\nВЫБРАННЫЙ ПРОЕКТ: {project}\nБАЗА:\n{context[:16000]}"
    content = [{"type": "text", "text": user_text}]
    for i, img in enumerate(selected, 1):
        content += [{"type": "text", "text": f"Изображение {i}: {img['name']}\nОписание: {img.get('description','')}\nКонтекст: {img.get('context','')[:1800]}"}, {"type": "image_url", "image_url": {"url": _data_url(img['path'])}}]

    payload = {"model": TOKENBOM_MODEL, "messages":[{"role":"system","content":system},{"role":"user","content":content}], "max_tokens":1200}
    started = asyncio.get_running_loop().time()
    answer = await asyncio.to_thread(_tokenbom, payload, key, 90)
    print(f"[Efin AI] TokenBom ответил за {asyncio.get_running_loop().time()-started:.2f} сек.")
    send = "[SEND_IMAGES]" in answer
    return {"answer": answer.replace("[SEND_IMAGES]", "").strip(), "images": selected if send else []}
