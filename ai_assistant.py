import asyncio
import base64
import io
import json
import os
import re
import urllib.error
import urllib.request
from pathlib import Path

from knowledge import get_projects, get_project_files

TOKENBOM_API_URL = "https://tokenbom.com/v1/chat/completions"
TOKENBOM_MODEL = "gpt-5.6-luna"
MEDIA_DIR = Path("data/knowledge_media")
MEDIA_DIR.mkdir(parents=True, exist_ok=True)

SYSTEM_PROMPT = """Ты — Efin AI, внутренний помощник компании EFIN.
Отвечай на русском языке, понятно и по делу.

БАЗА ЗНАНИЙ EFIN — главный источник.

Правила:
1. Используй материалы ВСЕХ файлов выбранного проекта. В проекте может быть до 10 файлов.
2. Не выдумывай внутренние правила, тарифы, выплаты, инструкции и условия.
3. Не смешивай материалы разных проектов.
4. Если пользователь просит фото, изображение, скриншот или пример, отправляй только реально подходящее изображение.
5. Наличие картинки в памятке само по себе НЕ означает, что она подходит.
6. Если подходящего изображения нет, не добавляй [SEND_IMAGES].
7. Если подходящее изображение есть, добавь в конце ответа ровно [SEND_IMAGES].
8. Не утверждай, что на картинке есть клиент, человек, карта, документ или другой объект, если это нельзя подтвердить.
9. Если информации нет в базе, прямо скажи об этом.
"""

_CACHE = {}


def _decode(data):
    for encoding in ("utf-8", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def _media_dir(project_id, file_id):
    directory = MEDIA_DIR / str(project_id) / str(file_id)
    directory.mkdir(parents=True, exist_ok=True)
    return directory


def _save_image(project_id, file_id, index, image_bytes, suffix, name, context="", page=None):
    suffix = (suffix or ".png").lower()
    if not suffix.startswith("."):
        suffix = "." + suffix
    path = _media_dir(project_id, file_id) / f"{index:03d}{suffix}"
    path.write_bytes(image_bytes)
    return {
        "path": str(path),
        "name": name or path.name,
        "context": context or "",
        "page": page,
        "description": "",
        "file_id": file_id,
    }


def _extract_pdf(data, project_id, file_id):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    text_parts = []
    images = []
    image_index = 0

    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            text_parts.append(f"Страница {page_number}:\n{text}")

        try:
            for image in page.images:
                image_index += 1
                name = image.name or f"page_{page_number}_{image_index}.png"
                suffix = Path(name).suffix.lower() or ".png"
                images.append(
                    _save_image(
                        project_id,
                        file_id,
                        image_index,
                        image.data,
                        suffix,
                        f"Страница {page_number} — {name}",
                        text,
                        page_number,
                    )
                )
        except Exception as exc:
            print(f"[Efin AI] Ошибка извлечения изображений PDF, страница {page_number}: {exc}")

    return "\n\n".join(text_parts), images


def _extract_pptx(data, project_id, file_id):
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    text_parts = []
    images = []
    image_index = 0

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            try:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    slide_text.append(shape.text.strip())
            except Exception:
                pass

        context = "\n".join(slide_text).strip()
        if context:
            text_parts.append(f"Слайд {slide_number}:\n{context}")

        for shape in slide.shapes:
            if getattr(shape, "shape_type", None) != 13:
                continue
            try:
                image_index += 1
                ext = "." + (shape.image.ext or "png").lower().lstrip(".")
                images.append(
                    _save_image(
                        project_id,
                        file_id,
                        image_index,
                        shape.image.blob,
                        ext,
                        f"Слайд {slide_number} — изображение {image_index}",
                        context,
                        slide_number,
                    )
                )
            except Exception as exc:
                print(f"[Efin AI] Ошибка извлечения изображения PPTX: {exc}")

    return "\n\n".join(text_parts), images


def _extract(data, file_name, mime_type, project_id, file_id):
    name = (file_name or "").lower()
    mime = (mime_type or "").lower()

    if name.endswith(".txt") or "text/plain" in mime:
        return _decode(data), []

    if name.endswith(".pdf") or "application/pdf" in mime:
        return _extract_pdf(data, project_id, file_id)

    if name.endswith(".pptx") or "presentationml.presentation" in mime:
        return _extract_pptx(data, project_id, file_id)

    if name.endswith(".docx") or "wordprocessingml.document" in mime:
        from docx import Document
        document = Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs if p.text.strip()), []

    if name.endswith((".jpg", ".jpeg", ".png", ".webp")) or mime.startswith("image/"):
        suffix = Path(name).suffix.lower() or ".jpg"
        return (
            "Изображение из базы знаний EFIN.",
            [_save_image(project_id, file_id, 1, data, suffix, file_name or "изображение")],
        )

    return _decode(data), []


def _image_data_url(path):
    data = Path(path).read_bytes()
    mime = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
    }.get(Path(path).suffix.lower(), "image/jpeg")
    return f"data:{mime};base64,{base64.b64encode(data).decode('ascii')}"


def _request(payload, api_key, timeout=60):
    request = urllib.request.Request(
        TOKENBOM_API_URL,
        data=json.dumps(payload).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"TokenBom API HTTP {exc.code}: {body[:1000]}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Не удалось подключиться к TokenBom API: {exc.reason}") from exc


def _extract_answer(data):
    try:
        return str(data["choices"][0]["message"]["content"] or "").strip()
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("TokenBom API вернул неожиданный формат ответа") from exc


async def _download_telegram_file(bot, file_id, attempts=2):
    last_error = None
    for attempt in range(1, attempts + 1):
        try:
            telegram_file = await bot.get_file(
                file_id,
                read_timeout=45,
                connect_timeout=15,
                pool_timeout=30,
            )
            data = await telegram_file.download_as_bytearray(
                read_timeout=90,
                connect_timeout=15,
                pool_timeout=30,
            )
            return bytes(data)
        except Exception as exc:
            last_error = exc
            print(f"[Efin AI] Ошибка загрузки файла, попытка {attempt}/{attempts}: {exc}")
            if attempt < attempts:
                await asyncio.sleep(1)
    raise last_error


async def _load_one_file(bot, project, file_row):
    cache_key = (project["id"], file_row["id"], file_row["file_id"])
    if cache_key in _CACHE:
        return _CACHE[cache_key]

    print(f"[Efin AI] Загружаю: {file_row['file_name']}")
    data = await _download_telegram_file(bot, file_row["file_id"])
    text, images = _extract(
        data,
        file_row["file_name"],
        file_row["mime_type"],
        project["id"],
        file_row["id"],
    )

    result = {
        "file_id": file_row["file_id"],
        "file_name": file_row["file_name"],
        "text": text.strip(),
        "images": images,
    }
    _CACHE[cache_key] = result
    print(
        f"[Efin AI] Готово: {file_row['file_name']} | "
        f"текст={len(result['text'])} | фото={len(images)}"
    )
    return result


async def _project_data(bot, project):
    files = get_project_files(project["id"])
    print(f"[Efin AI] Проект «{project['name']}»: файлов {len(files)}")

    # Файлы скачиваются параллельно: один медленный PDF не блокирует остальные.
    results = await asyncio.gather(
        *[_load_one_file(bot, project, row) for row in files],
        return_exceptions=True,
    )

    all_text = []
    all_images = []

    for row, result in zip(files, results):
        if isinstance(result, Exception):
            print(f"[Efin AI] Не удалось прочитать файл {row['file_name']}: {result}")
            continue
        all_text.append(f"=== ФАЙЛ: {row['file_name']} ===\n{result['text']}")
        all_images.extend(result["images"])

    print(
        f"[Efin AI] Проект «{project['name']}»: "
        f"текстов={len(all_text)}, изображений={len(all_images)}"
    )
    return {"text": "\n\n".join(all_text), "images": all_images}


def _words(text):
    return set(re.findall(r"[а-яa-z0-9ё-]{3,}", (text or "").lower()))


def _relevant(text, query, limit=9000):
    if not text:
        return ""

    paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    words = _words(query)
    scored = []

    for index, paragraph in enumerate(paragraphs):
        low = paragraph.lower()
        score = sum(1 for word in words if word in low)
        scored.append((score, index, paragraph))

    scored.sort(key=lambda item: (-item[0], item[1]))
    selected = [item for item in scored if item[0] > 0][:30] or scored[:10]

    result = []
    total = 0
    for _, index, paragraph in selected:
        block = f"Фрагмент {index + 1}:\n{paragraph}"
        if total + len(block) > limit:
            break
        result.append(block)
        total += len(block)

    return "\n\n".join(result)


def _visual_query(query):
    text = (query or "").lower()
    return any(
        token in text
        for token in (
            "фото", "фотограф", "фотографию", "фотография", "сфот",
            "изображ", "картин", "пример", "скрин", "скриншот", "снимок",
            "покажи", "пришли", "показывай", "как выглядит", "как правильно",
            "камера", "визуал",
        )
    )


def _score_image(image, query):
    words = _words(query)
    searchable = " ".join(
        (
            image.get("name", ""),
            image.get("context", ""),
            image.get("description", ""),
        )
    ).lower()

    if not words or not searchable:
        return 0

    score = 0
    for word in words:
        if word in searchable:
            score += 1
            if len(word) >= 6:
                score += 2

    pairs = (
        (("клиент", "клиента", "клиентом"), ("фото", "фотография", "фотографию")),
        (("пример", "образец"), ("фото", "фотография", "фотографию")),
        (("карта", "карту"), ("банков", "банковская", "банковскую")),
    )
    for left, right in pairs:
        if any(w in words for w in left) and any(w in searchable for w in right):
            score += 3

    return score


def _candidate_images(images, query, limit=6):
    scored = [(_score_image(image, query), index, image) for index, image in enumerate(images)]
    scored.sort(key=lambda item: (-item[0], item[1]))
    return [item[2] for item in scored[:limit] if item[0] > 0]


def _analyse_image_batch_sync(images, api_key):
    if not images or not api_key:
        return []

    content = [{
        "type": "text",
        "text": (
            "Проанализируй изображения из внутренней базы EFIN. "
            "Для каждого изображения отдельно дай короткое точное описание. "
            "Не придумывай ничего. Определи, если это действительно видно: "
            "человек/клиент, фото клиента, карта, документ, телефон, экран, QR-код, "
            "штрихкод, логотип или другой объект. "
            "Ответь строго строками IMAGE 1: ... IMAGE 2: ... и т.д."
        ),
    }]

    for index, image in enumerate(images, start=1):
        content.append({
            "type": "text",
            "text": f"IMAGE {index}. Имя: {image.get('name', '')}. Контекст: {image.get('context', '')[:1200]}",
        })
        content.append({
            "type": "image_url",
            "image_url": {"url": _image_data_url(image["path"])},
        })

    payload = {
        "model": TOKENBOM_MODEL,
        "messages": [
            {"role": "system", "content": "Ты быстро классифицируешь изображения внутренней базы. Только подтверждаемые факты."},
            {"role": "user", "content": content},
        ],
        "max_tokens": 900,
    }

    try:
        data = _request(payload, api_key, timeout=75)
        answer = _extract_answer(data)
        descriptions = {}
        for line in answer.splitlines():
            match = re.match(r"IMAGE\s+(\d+)\s*:\s*(.*)", line.strip(), re.I)
            if match:
                descriptions[int(match.group(1))] = match.group(2).strip()
        return [descriptions.get(index, "") for index in range(1, len(images) + 1)]
    except Exception as exc:
        print(f"[Efin AI] Ошибка группового анализа изображений: {exc}")
        return [""] * len(images)


async def _analyse_images_batch(images, api_key):
    if not images or not api_key:
        return

    # Один запрос анализирует до 6 картинок вместо отдельного запроса на каждую.
    for start in range(0, len(images), 6):
        batch = images[start:start + 6]
        descriptions = await asyncio.to_thread(_analyse_image_batch_sync, batch, api_key)
        for image, description in zip(batch, descriptions):
            if description:
                image["description"] = description


async def build_knowledge_context(bot, user_text):
    projects = get_projects()
    if not projects:
        return "База знаний EFIN пока пуста.", []

    visual = _visual_query(user_text)
    parts = []
    selected_images = []
    api_key = os.getenv("TOKENBOOM_API_KEY")

    for project in projects:
        data = await _project_data(bot, project)
        relevant = _relevant(data["text"], user_text)

        if relevant:
            parts.append(f"=== ПРОЕКТ: {project['name']} ===\n{relevant}")

        if visual and data["images"]:
            # Сначала быстрый поиск без API.
            candidates = _candidate_images(data["images"], user_text, limit=4)

            # Только если быстрый поиск не нашёл ничего — запускаем vision.
            if not candidates and api_key:
                print(
                    f"[Efin AI] В проекте «{project['name']}» быстрый поиск не нашёл фото. "
                    "Запускаю групповой анализ изображений..."
                )
                await _analyse_images_batch(data["images"], api_key)
                candidates = _candidate_images(data["images"], user_text, limit=4)

            selected_images.extend(candidates)

    unique = []
    seen = set()
    for image in selected_images:
        if image["path"] in seen:
            continue
        seen.add(image["path"])
        unique.append(image)

    return (
        "\n\n".join(parts)[:14000]
        or "В базе знаний EFIN нет найденного текстового фрагмента по этому вопросу.",
        unique[:6],
    )


async def ask_efin_ai(bot, user_text):
    api_key = os.getenv("TOKENBOOM_API_KEY")
    if not api_key:
        raise RuntimeError("Не задана переменная окружения TOKENBOOM_API_KEY")

    knowledge, images = await build_knowledge_context(bot, user_text)
    print(
        f"[Efin AI] Контекст базы: {len(knowledge)} символов; "
        f"релевантных изображений: {len(images)}"
    )

    system_text = SYSTEM_PROMPT + "\n\nБАЗА ЗНАНИЙ EFIN:\n" + knowledge
    user_content = [{"type": "text", "text": user_text}]

    for index, image in enumerate(images, start=1):
        user_content.append({
            "type": "text",
            "text": (
                f"Изображение {index}: {image['name']}\n"
                f"Описание: {image.get('description', '')[:3000]}\n"
                f"Контекст: {image.get('context', '')[:1800]}"
            ),
        })
        user_content.append({
            "type": "image_url",
            "image_url": {"url": _image_data_url(image["path"])},
        })

    payload = {
        "model": TOKENBOM_MODEL,
        "messages": [
            {"role": "system", "content": system_text},
            {"role": "user", "content": user_content},
        ],
        "max_tokens": 1200,
    }

    started = asyncio.get_running_loop().time()
    data = await asyncio.to_thread(_request, payload, api_key, 60)
    elapsed = asyncio.get_running_loop().time() - started
    print(f"[Efin AI] TokenBom ответил за {elapsed:.2f} сек.")

    answer = _extract_answer(data)
    if not answer:
        raise RuntimeError("TokenBom API вернул пустой ответ")

    send_images = "[SEND_IMAGES]" in answer
    clean_answer = answer.replace("[SEND_IMAGES]", "").strip()

    return {
        "answer": clean_answer,
        "images": images if send_images else [],
    }
